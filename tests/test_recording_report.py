import csv
import importlib.util
import io
import json
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

spec = importlib.util.spec_from_file_location("recording_report", Path(__file__).parents[1]/"backend/src/recording_report.py")
report = importlib.util.module_from_spec(spec)
spec.loader.exec_module(report)


class Reports(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.name = "groundstation_recording_2026-09-20_00-00-00_000.db"
        self.path = self.root/self.name
        with sqlite3.connect(self.path) as db:
            db.execute("CREATE TABLE telemetry(id INTEGER PRIMARY KEY,timestamp_ms INTEGER,source_timestamp_ms INTEGER,sender_id TEXT,data_type TEXT,values_json TEXT,payload_json TEXT)")
            for i in range(10):
                for kind, value in [("KG1000", i), ("LOADCELL_WEIGHT_KG",2*i-3)]:
                    db.execute("INSERT INTO telemetry VALUES(NULL,?,?,?,?,?,?)",(1000+i*100, i,"DAQ",kind,json.dumps([value]),"[]"))
        self.request={"root":str(self.root),"recording":self.name}

    def tearDown(self):
        self.tmp.cleanup()

    def test_preview_range_channels_and_regression(self):
        preview=json.loads(report.run(self.request))
        self.assertEqual(preview["recorded_rows"],20)
        cal=next(c for c in preview["channels"] if "LOADCELL_WEIGHT_KG" in c["key"])
        self.assertEqual(cal["regression"]["slope"],2)
        self.assertEqual(cal["regression"]["intercept"],-3)
        self.assertEqual(cal["regression"]["status"],"inferred_from_recorded_pairs")
        data=report.run(dict(self.request,format="csv",start_ms=1200,end_ms=1600,channels=[cal["key"]]))
        rows=list(csv.DictReader(io.StringIO(data.decode())))
        self.assertEqual(len(rows),4)
        self.assertEqual([float(r["value"]) for r in rows],[1,3,5,7])
        self.assertTrue(all(r["regression_json"] for r in rows))

    def test_invalid_selection_and_symlink_rejected(self):
        for extra in [{"recording":"../secrets.db"},{"start_ms":3,"end_ms":2},{"channels":[]},{"channels":["missing"]}]:
            with self.assertRaises(ValueError):report.run(dict(self.request,**extra))
        link=self.root/"groundstation_recording_1.db"
        link.symlink_to(self.path)
        with self.assertRaises(ValueError):report.run(dict(self.request,recording=link.name))

    def test_actual_calibration_metadata_takes_precedence_and_tracks_changes(self):
        with sqlite3.connect(self.path) as db:
            db.execute("CREATE TABLE calibration_history(timestamp_ms INTEGER,sender_id TEXT,sensor_id TEXT,config_json TEXT)")
            for t,m in [(1000,2),(1400,3),(1700,2)]:
                db.execute("INSERT INTO calibration_history VALUES(?,?,?,?)",(t,"DAQ","KG1000",json.dumps({"calibration":{"ch1":{"m":m,"b":-3},"ch1_fit":{"type":"linear"},"ch1_zero_raw":None}})))
        preview=json.loads(report.run(self.request))
        cal=next(c for c in preview["channels"] if "LOADCELL_WEIGHT_KG" in c["key"])
        self.assertEqual(cal["regression"]["status"],"recorded_calibration")
        self.assertEqual(len(cal["regression"]["epochs"]),3)

    def test_preview_preserves_spikes_and_gaps(self):
        samples=[(i,1000 if i==451 else 0,None,"r",i) for i in range(5000)]
        samples[300]=(300,None,None,"r",300)
        points=report.preview_points(samples)
        self.assertLessEqual(len(points),600)
        self.assertTrue(any(y==1000 for _,y in points))
        self.assertTrue(any(y is None for _,y in points))

    def test_fill_percent_metadata_preserves_selected_source_and_absolute_mapping(self):
        with sqlite3.connect(self.path) as db:
            db.execute("CREATE TABLE calibration_history(timestamp_ms INTEGER,sender_id TEXT,sensor_id TEXT,config_json TEXT)")
            for t, source, sensor in [(1000,"kg50","KG50"),(1400,"kg1000_absolute","KG1000")]:
                snapshot={"calibration":{"ch1":{"m":2,"b":1},"extra_channels":{"kg50":{"linear":{"m":3,"b":4}}}},
                          "fill_targets":{"fill_source":source,"nitrous":{"target_mass_kg":10}}}
                db.execute("INSERT INTO calibration_history VALUES(?,?,?,?)",(t,"DAQ",sensor,json.dumps(snapshot)))
                db.execute("INSERT INTO telemetry VALUES(NULL,?,?,?,?,?,?)",(t,t,"DAQ","LOADCELL_FILL_PERCENT","[50]","[]"))
        preview=json.loads(report.run(self.request))
        cal=next(c for c in preview["channels"] if "LOADCELL_FILL_PERCENT" in c["key"])
        epochs=cal["regression"]["epochs"]
        self.assertEqual([e["fill_source"] for e in epochs],["KG50","KG1000"])
        self.assertNotIn("abs(",epochs[0]["postprocess"])
        self.assertIn("abs(calibrated_kg)",epochs[1]["postprocess"])

    def test_real_excel_has_chart_and_regression(self):
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(report.run(dict(self.request,format="xlsx"))))
        self.assertIn("Regression metadata",book.sheetnames)
        self.assertEqual(len(book["Channel 1"]._charts),1)
        self.assertEqual(book["Channel 1"].max_row,11)
        book.close()

    def test_inconsistent_historical_mapping_is_not_misrepresented(self):
        with sqlite3.connect(self.path) as db:
            db.execute("UPDATE telemetry SET values_json='[500]' WHERE id=10")
        preview=json.loads(report.run(self.request))
        cal=next(c for c in preview["channels"] if "LOADCELL_WEIGHT_KG" in c["key"])
        self.assertEqual(cal["regression"]["status"],"unavailable")

    def test_all_recordings_and_limits(self):
        second=self.root/"groundstation_recording_2.db"
        second.write_bytes(self.path.read_bytes())
        preview=json.loads(report.run(dict(self.request,recording="")))
        self.assertEqual(preview["recorded_rows"],40)
        old=report.MAX_ROWS
        try:
            report.MAX_ROWS=5
            with self.assertRaisesRegex(ValueError,"no rows were silently dropped"):
                report.run(self.request)
        finally:
            report.MAX_ROWS=old

    def test_real_pdf_contains_graph_data_and_regression(self):
        from pypdf import PdfReader
        pdf=PdfReader(io.BytesIO(report.run(dict(self.request,format="pdf"))))
        self.assertEqual(len(pdf.pages),6)
        text="\n".join(page.extract_text() for page in pdf.pages)
        self.assertIn("slope",text)
        self.assertIn("Received ms",text)


if __name__ == "__main__": unittest.main()
