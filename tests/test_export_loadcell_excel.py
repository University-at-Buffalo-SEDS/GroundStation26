import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from export_loadcell_excel import Sample, export_excel, extrema, parse_time, read_recording, recalculate_kg, select_window


class RecordingExportTests(unittest.TestCase):
    def test_regression_absolute_values_and_saved_coefficients(self):
        raw = [Sample(1000 + i, i, "DAQ", str(i), value)
               for i, value in enumerate((1.0, 2.0, 3.0, None))]
        kg = [Sample(row.received_ms, row.source_ms, row.sender, row.row_id,
                     None if row.value is None else 10 * row.value - 25) for row in raw]
        streams, notes, fits = recalculate_kg({"KG1000": raw, "LOADCELL_WEIGHT_KG": kg})
        self.assertEqual(fits[0]["slope"], 10)
        self.assertEqual(fits[0]["intercept"], -25)
        self.assertEqual(fits[0]["r_squared"], 1)
        self.assertEqual([row.value for row in streams["ABSOLUTE_KG"]], [15, 5, 5, None])
        self.assertIs(streams["LOADCELL_WEIGHT_KG"], kg)
        low, high = extrema(streams["RECALCULATED_KG"])
        self.assertEqual((low.value, high.value), (-15, 5))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "regression.xlsx"
            export_excel(streams, path, Path("recording.csv"), 1000, 1001, notes=notes, fits=fits)
            book = load_workbook(path)
            # Regression includes all matched pairs; selected data includes only two.
            self.assertEqual(book["Calibration"]["C7"].value, 10)
            self.assertEqual(book["Calibration"]["D7"].value, -25)
            self.assertEqual(book["Calibration"]["E7"].value, 3)
            self.assertEqual(len(book["Calibration pairs 1"]._charts), 1)
            self.assertEqual(book["Loadcell raw"].max_row, 3)
            self.assertEqual(book["Min and max"]["C2"].value, 1)
            self.assertEqual(book["Min and max"]["E2"].value, 2)
            self.assertEqual(len(book["Charts"]._charts[0].series), 3)
            book.close()

    def test_inference_rejects_clipped_mapping_but_accepts_known_coefficients(self):
        raw = [Sample(i, i, "DAQ", str(i), float(i)) for i in range(5)]
        kg = [Sample(i, i, "DAQ", str(i), float(max(i, 2))) for i in range(5)]
        with self.assertRaisesRegex(ValueError, "not a consistent linear mapping"):
            recalculate_kg({"KG1000": raw, "LOADCELL_WEIGHT_KG": kg})
        streams, _, fits = recalculate_kg({"KG1000": raw, "LOADCELL_WEIGHT_KG": kg}, 1, 0)
        self.assertEqual(streams["RECALCULATED_KG"][0].value, 0)
        self.assertEqual(fits[0]["method"], "User supplied")
        with self.assertRaises(ValueError):
            recalculate_kg({"KG1000": raw}, float("nan"), 0)

    def test_read_preserves_nulls_and_duplicate_timestamps(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "recording.csv"
            with path.open("w", newline="") as output:
                writer = csv.writer(output)
                writer.writerow(("id", "received_timestamp_ms", "source_timestamp_ms",
                                 "sender_id", "data_type", "values_json"))
                writer.writerows([
                    (1, 2000, 10, "DAQ", "KG1000", "[1.5]"),
                    (2, 1000, 9, "DAQ", "KG1000", "[null]"),
                    (3, 2000, 11, "DAQ", "KG1000", "[2.5]"),
                    (4, 2000, 44, "VB", "FUEL_TANK_PRESSURE", "[3.5]"),
                    (5, 2000, 11, "DAQ", "KG50", "[99]"),
                ])
            streams = read_recording(path)
            self.assertEqual([row.value for row in streams["KG1000"]], [None, 1.5, 2.5])
            self.assertEqual(len(select_window(streams, 2000, 2000)["KG1000"]), 2)
            self.assertNotIn("KG50", streams)
            self.assertEqual(streams["FUEL_TANK_PRESSURE"][0].source_ms, 44)
            with self.assertRaises(ValueError):
                select_window(streams, 3000, 4000)
            with self.assertRaises(ValueError):
                select_window(streams, 2000, 1000)

    def test_excel_preserves_times_values_and_numeric_chart_axis(self):
        timestamp = 1789794290932
        streams = {
            "KG1000": [Sample(timestamp, 17744228, "=literal", "3", 0.015170760452747345)],
            "LOADCELL_WEIGHT_KG": [Sample(timestamp, 17744228, "DAQ", "4", 4.047607421875)],
            "FUEL_TANK_PRESSURE": [Sample(timestamp + 50, 17744480, "VB", "16", 466.6)],
            "PRESSURE_TRANSDUCER_CALIBRATED": [Sample(timestamp + 50, 17744480, "VB", "17", None)],
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "output.xlsx"
            self.assertEqual(export_excel(streams, path, Path("input.csv"), timestamp, timestamp + 50), 4)
            book = load_workbook(path)
            self.assertEqual(len(book["Charts"]._charts), 4)
            sheet = book["Loadcell raw"]
            self.assertEqual(sheet["A2"].value.microsecond, 932000)
            self.assertEqual(sheet["B2"].value, timestamp)
            self.assertEqual(sheet["C2"].value, 17744228)
            self.assertEqual(sheet["D2"].data_type, "s")
            self.assertAlmostEqual(sheet["F2"].value, streams["KG1000"][0].value)
            self.assertIsNone(book["PT calibrated"]["F2"].value)
            chart = book["Charts"]._charts[0]
            self.assertEqual(chart.series[0].xVal.numRef.f, "'Loadcell raw'!$A$2")
            book.close()
            with self.assertRaises(FileExistsError):
                export_excel(streams, path, Path("input.csv"), timestamp, timestamp + 50)

    def test_timestamp_offsets(self):
        self.assertEqual(parse_time("2026-09-19T05:04:50.932-04:00"),
                         parse_time("2026-09-19T09:04:50.932Z"))
        self.assertEqual(parse_time("2026-09-19T09:04:50.932"), 1789808690932)


if __name__ == "__main__":
    unittest.main()
