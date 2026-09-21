"""Headless recording preview/report worker. JSON request on stdin, result on stdout.

CSV/preview need only Python's standard library. XLSX uses openpyxl; PDF uses
matplotlib's Agg/PdfPages backend, never a display server or browser window.
"""
import csv
import io
import json
import math
import re
import sqlite3
import sys
from collections import defaultdict
from pathlib import Path

MAX_ROWS = 200_000
MAX_SERIES = 128
RAW_TYPES = {"LOADCELL_WEIGHT_KG": "KG1000", "LOADCELL_50_WEIGHT_KG": "KG50",
             "PRESSURE_TRANSDUCER_CALIBRATED": "FUEL_TANK_PRESSURE"}


def read_rows(request):
    root = Path(request["root"]).resolve(strict=True)
    recording = request.get("recording", "")
    paths = []
    for path in root.iterdir():
        if not re.fullmatch(r"groundstation_recording_[0-9_-]+\.db", path.name):
            continue
        if recording and path.name != recording:
            continue
        if path.is_symlink() or not path.is_file() or path.resolve().parent != root:
            continue
        paths.append(path)
    if recording and not paths:
        raise ValueError("Recording not found. Refresh the recording list.")
    start, end = request.get("start_ms"), request.get("end_ms")
    if start is not None and end is not None and end <= start:
        raise ValueError("End time must be after start time.")
    result, history = [], defaultdict(list)
    for path in sorted(paths):
        with sqlite3.connect(path.as_uri() + "?mode=ro", uri=True, timeout=5) as db:
            db.execute("BEGIN")
            if db.execute("SELECT 1 FROM sqlite_master WHERE name='calibration_history'").fetchone():
                for t, sender, sensor, config in db.execute("SELECT timestamp_ms,sender_id,sensor_id,config_json FROM calibration_history ORDER BY timestamp_ms"):
                    entries = history[(path.name, sender, sensor)]
                    snapshot = json.loads(config)
                    if not entries or entries[-1][1] != snapshot:
                        entries.append((t, snapshot))
            upper = db.execute("SELECT COALESCE(MAX(id),0) FROM telemetry").fetchone()[0]
            sql = "SELECT id,timestamp_ms,source_timestamp_ms,sender_id,data_type,values_json,payload_json FROM telemetry WHERE id<=?"
            args = [upper]
            if start is not None:
                sql += " AND timestamp_ms>=?"
                args.append(start)
            if end is not None:
                sql += " AND timestamp_ms<?"
                args.append(end)
            for row in db.execute(sql + " ORDER BY id", args):
                result.append((path.name,) + row)
                if len(result) > MAX_ROWS:
                    raise ValueError("Selection exceeds 200,000 recorded rows. Choose a recording or narrower time window; no rows were silently dropped.")
    if not result:
        raise ValueError("No recorded samples in this selection.")
    return sorted(result, key=lambda row: (row[2], row[0], row[1])), history


def channels(rows):
    result = defaultdict(list)
    count = 0
    for recording, row_id, received, source, sender, kind, values, payload in rows:
        for index, value in enumerate(json.loads(values or "[]")):
            count += 1
            if count > MAX_ROWS:
                raise ValueError("Selection exceeds 200,000 channel samples. Choose a narrower time window; no data was silently dropped.")
            if value is not None and (not isinstance(value, (float, int)) or not math.isfinite(value)):
                value = None
            key = json.dumps([sender or "", kind, index], separators=(",", ":"))
            result[key].append((received, value, source, recording, row_id))
            if len(result) > MAX_SERIES:
                raise ValueError("Too many channels for one report. Select a smaller recording.")
    return dict(result)


def regression(key, samples, streams, history=None):
    sender, kind, index = json.loads(key)
    sensor = {**RAW_TYPES, "LOADCELL_FILL_PERCENT": "KG50"}.get(kind)
    if sensor == "FUEL_TANK_PRESSURE":
        sensor = "IADC"
    if sensor and history:
        import bisect
        epochs, missing = {}, 0
        indexes = {}
        for t, _, _, recording, _ in samples:
            if recording not in indexes:
                records = history.get((recording, sender, sensor), [])
                if kind == "LOADCELL_FILL_PERCENT":
                    records = sorted((entry for source in ("KG50", "KG1000")
                        for entry in history.get((recording, sender, source), [])
                        if source == ("KG1000" if entry[1].get("fill_targets", {}).get("fill_source") == "kg1000_absolute" else "KG50")), key=lambda entry: entry[0])
                indexes[recording] = (records, [entry[0] for entry in records])
            records, stamps = indexes[recording]
            i = bisect.bisect_right(stamps, t)-1
            if i < 0:
                missing += 1
                continue
            stamp, snapshot = records[i]
            cfg = snapshot["calibration"]
            sample_sensor = sensor
            if kind == "LOADCELL_FILL_PERCENT":
                sample_sensor = "KG1000" if snapshot.get("fill_targets", {}).get("fill_source") == "kg1000_absolute" else "KG50"
            channel_name = {"KG1000":"ch1", "KG50":"kg50", "IADC":"iadc"}[sample_sensor]
            if channel_name == "kg50":
                channel = cfg.get("extra_channels", {}).get("kg50", {"linear":{"m":1,"b":0}})
            else:
                channel = {"linear":cfg[channel_name], "fit":cfg.get(channel_name+"_fit"),
                           "zero_raw":cfg.get(channel_name+"_zero_raw"),
                           "points":cfg.get("points_"+channel_name, [])}
            signature = json.dumps([recording, stamp, channel], sort_keys=True)
            entry = epochs.setdefault(signature, {"recording":recording,"first_sample_ms":t,"last_sample_ms":t,
                "channel":channel,"equation":"y = f(raw) - f(zero_raw) if tare exists, else f(raw); linear: f(x)=m*x+b; polyN: descending a..e coefficients evaluated at (x-x0)",
                "samples":0})
            entry["last_sample_ms"] = t
            entry["samples"] += 1
            if kind == "LOADCELL_FILL_PERCENT":
                entry["fill_targets"] = snapshot.get("fill_targets")
                entry["flight_state"] = snapshot.get("flight_state")
                entry["fill_source"] = sample_sensor
                entry["postprocess"] = "clamp(100 * " + ("abs(calibrated_kg)" if sample_sensor == "KG1000" else "calibrated_kg") + " / active_target_kg, 0, 100)"
        if epochs:
            return {"status":"recorded_calibration" if missing==0 else "partially_recorded_calibration",
                    "epochs":list(epochs.values()), "samples_without_metadata":missing}
    if kind not in RAW_TYPES:
        if "CALIBRAT" in kind or kind == "LOADCELL_FILL_PERCENT":
            return {"status": "unavailable", "reason": "No historical regression recorded for this derived channel."}
        return {"status": "not_applicable"}
    raw = streams.get(json.dumps([sender, RAW_TYPES[kind], index], separators=(",", ":")), [])
    lookup = {(t, src, recording): value for t, value, src, recording, _ in raw}
    pairs = [(lookup[(t, src, rec)], value) for t, value, src, rec, _ in samples
             if value is not None and lookup.get((t, src, rec)) is not None]
    if len(pairs) < 3:
        return {"status": "unavailable", "reason": "Insufficient matched raw/calibrated samples; current calibration is not substituted."}
    mx = sum(x for x, _ in pairs) / len(pairs)
    my = sum(y for _, y in pairs) / len(pairs)
    xx = sum((x-mx)**2 for x, _ in pairs)
    if xx < 1e-24:
        return {"status": "unavailable", "reason": "Raw samples have insufficient variation."}
    slope = sum((x-mx)*(y-my) for x, y in pairs) / xx
    intercept = my-slope*mx
    residual = max(abs(y-(slope*x+intercept)) for x, y in pairs)
    tolerance = max(1e-5, max(abs(y) for _, y in pairs)*2e-6)
    if residual > tolerance:
        return {"status": "unavailable", "reason": "Samples do not support one linear regression (possibly calibration changes or a nonlinear fit). Select a narrower interval."}
    return {"status": "inferred_from_recorded_pairs", "equation": "y = slope * x + intercept",
            "raw_type": RAW_TYPES[kind], "slope": slope, "intercept": intercept,
            "pair_count": len(pairs), "max_residual": residual,
            "note": "Reconstructed mapping, not a saved original calibration."}


def preview_points(samples, limit=600):
    # Keep extrema rather than averaging away brief pressure/loadcell peaks.
    width = max(1, math.ceil(len(samples)/(limit/5)))
    points = []
    for offset in range(0, len(samples), width):
        group = samples[offset:offset+width]
        valid = [(i, s) for i, s in enumerate(group) if s[1] is not None]
        indexes = {0, len(group)-1}
        missing = next((i for i,s in enumerate(group) if s[1] is None), None)
        if missing is not None:
            indexes.add(missing)
        if valid:
            indexes.add(min(valid, key=lambda p:p[1][1])[0])
            indexes.add(max(valid, key=lambda p:p[1][1])[0])
        points.extend([[group[i][0], group[i][1]] for i in sorted(indexes)])
    return points


def selected_streams(request, streams):
    selected = request.get("channels")
    if selected is None:
        return streams
    if not isinstance(selected, list) or not selected:
        raise ValueError("Select at least one channel, or choose all data.")
    unknown = set(selected)-set(streams)
    if unknown:
        raise ValueError("A selected channel has no data in this interval. Refresh the preview or change the selection.")
    return {key: streams[key] for key in selected}


def export_csv(streams, regressions):
    if sum(len(json.dumps(regressions[key]))*len(samples) for key,samples in streams.items()) > 48_000_000:
        raise ValueError("Regression metadata makes this CSV too large. Narrow the selection or use Excel, which stores regressions separately.")
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    writer.writerow(["recording", "row_id", "received_timestamp_ms", "source_timestamp_ms",
                     "sender", "data_type", "channel", "value", "regression_json"])
    for key, samples in streams.items():
        sender, kind, index = json.loads(key)
        metadata = json.dumps(regressions[key], separators=(",", ":"))
        for t, value, source, recording, row_id in samples:
            # Spreadsheet formula protection for untrusted textual identities.
            safe = lambda s: "'"+s if s.startswith(("=", "+", "-", "@")) else s
            writer.writerow([recording, row_id, t, source, safe(sender), safe(kind), index, value, metadata])
    return out.getvalue().encode()


def export_xlsx(streams, regressions):
    from openpyxl import Workbook
    from openpyxl.chart import ScatterChart, Reference, Series
    book = Workbook()
    meta = book.active
    meta.title = "Regression metadata"
    meta.append(["Channel", "Regression / provenance"])
    for number, (key, samples) in enumerate(streams.items(), 1):
        metadata = json.dumps(regressions[key])
        for offset in range(0, len(metadata), 30000):
            meta.append([key, metadata[offset:offset+30000]])
            meta.cell(meta.max_row, 1).data_type = "s"
        sheet = book.create_sheet(f"Channel {number}")
        sheet.append(["Received ms", "Value", "Source ms", "Recording", "Row ID", "Channel"])
        for row in samples:
            t, value, source, recording, row_id = row
            sheet.append([t, value, source, recording, row_id, key])
            sheet.cell(sheet.max_row, 6).data_type = "s"
        chart = ScatterChart()
        chart.title = key
        chart.x_axis.title = "Received timestamp (ms UTC)"
        chart.y_axis.title = "Recorded value"
        chart.series.append(Series(Reference(sheet, min_col=2, min_row=2, max_row=sheet.max_row),
                                   Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row)))
        chart.width, chart.height = 25, 12
        sheet.add_chart(chart, "H2")
        sheet.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def export_pdf(streams, regressions):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    import textwrap
    # A raw-data PDF can be enormous. Reject instead of truncating the appendix.
    if sum(len(v) for v in streams.values()) > 20_000:
        raise ValueError("PDF raw-data appendix exceeds 20,000 samples. Narrow the selection or use CSV/Excel.")
    out = io.BytesIO()
    with PdfPages(out) as pdf:
        for key, samples in streams.items():
            fig, ax = plt.subplots(figsize=(11.7, 8.3))
            start = samples[0][0]
            ax.plot([(s[0]-start)/1000 for s in samples], [s[1] for s in samples], linewidth=.7)
            ax.set(title=key, xlabel=f"Seconds after receive UTC epoch {start} ms", ylabel="Recorded value")
            ax.grid(True, alpha=.3)
            pdf.savefig(fig)
            plt.close(fig)
            lines = textwrap.wrap(json.dumps(regressions[key]), 110)
            for offset in range(0,len(lines),48):
                fig, ax = plt.subplots(figsize=(11.7,8.3))
                ax.axis("off")
                ax.set_title(key+" — regression / provenance",fontsize=9)
                ax.text(0,1,"\n".join(lines[offset:offset+48]),va="top",fontsize=8, family="monospace")
                pdf.savefig(fig)
                plt.close(fig)
            for offset in range(0, len(samples), 35):
                fig, ax = plt.subplots(figsize=(11.7, 8.3))
                ax.axis("off")
                ax.set_title(f"{key} — recorded data {offset+1}–{min(offset+35,len(samples))}", fontsize=9)
                cells = [[str(v) if v is not None else "" for v in row] for row in samples[offset:offset+35]]
                table = ax.table(cellText=cells, colLabels=["Received ms", "Value", "Source ms", "Recording", "Row ID"], loc="center")
                table.auto_set_font_size(False)
                table.set_fontsize(5)
                table.scale(1, 1.35)
                pdf.savefig(fig)
                plt.close(fig)
    return out.getvalue()


def run(request):
    rows, history = read_rows(request)
    all_streams = channels(rows)
    streams = selected_streams(request, all_streams)
    metadata = {key: regression(key, samples, all_streams, history) for key, samples in streams.items()}
    fmt = request.get("format", "preview")
    if fmt == "preview":
        return json.dumps({"start_ms": rows[0][2], "end_ms": rows[-1][2]+1,
            "recorded_rows": len(rows), "channels": [{"key": key, "count": len(samples),
                "points": preview_points(samples), "regression": metadata[key]}
                for key, samples in streams.items()]}).encode()
    if fmt == "csv":
        return export_csv(streams, metadata)
    if fmt == "xlsx":
        return export_xlsx(streams, metadata)
    if fmt == "pdf":
        return export_pdf(streams, metadata)
    raise ValueError("Choose CSV, Excel or PDF.")


if __name__ == "__main__":
    try:
        sys.stdout.buffer.write(run(json.load(sys.stdin)))
    except ImportError as exc:
        sys.stderr.write(f"Report dependency unavailable: {exc}. Install openpyxl and matplotlib in the backend report Python environment; CSV/preview remain available.")
        sys.exit(2)
    except Exception as exc:
        sys.stderr.write(str(exc))
        sys.exit(1)
