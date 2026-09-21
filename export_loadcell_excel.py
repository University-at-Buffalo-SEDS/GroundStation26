#!/usr/bin/env python3
"""Select a recording time window and export load-cell/PT data and Excel charts.

Install: python3 -m pip install matplotlib openpyxl
Run: python3 export_loadcell_excel.py /path/to/recording.csv
Drag across any plot, then click Export selection. Closing the window cancels.
For unattended exports use --start/--end (ISO timestamps; UTC if no offset),
or --all. Received timestamps provide the shared clock; device source timestamps
are retained separately because different senders can have different clocks.
By default, infer the recording's linear raw-to-kg mapping and plot its absolute
value. Recorded and recalculated signed kilograms are also exported. Use
--kg-mode signed/recorded to change the plotted kilograms, or supply a known
mapping with --kg-slope and --kg-intercept (kg = slope * raw + intercept).
"""

import argparse
import csv
import json
import math
import sys
from collections import defaultdict, deque
from dataclasses import dataclass, replace
from datetime import datetime, timezone
from pathlib import Path


# Raw and recorded streams remain intact alongside explicitly derived streams.
CHANNELS = {
    "KG1000": ("Loadcell raw", "Raw"),
    "LOADCELL_WEIGHT_KG": ("Loadcell kg", "kg"),
    "FUEL_TANK_PRESSURE": ("PT raw", "Raw"),
    "IADC": ("PT raw IADC", "Raw"),
    "PRESSURE_TRANSDUCER_CALIBRATED": ("PT calibrated", "psi"),
    "RECALCULATED_KG": ("Loadcell recalculated kg", "kg"),
    "ABSOLUTE_KG": ("Loadcell absolute kg", "kg"),
}


@dataclass(frozen=True)
class Sample:
    received_ms: int
    source_ms: int | None
    sender: str
    row_id: str
    value: float | None

    @property
    def time(self):
        return datetime.fromtimestamp(self.received_ms / 1000, timezone.utc)


def read_recording(path):
    streams = {key: [] for key in CHANNELS}
    with path.open(newline="", encoding="utf-8-sig") as source:
        reader = csv.DictReader(source)
        required = {"received_timestamp_ms", "source_timestamp_ms", "sender_id",
                    "data_type", "values_json", "id"}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"CSV missing columns: {', '.join(sorted(missing))}")
        for line, row in enumerate(reader, 2):
            kind = row["data_type"]
            if kind not in streams:
                continue
            try:
                values = json.loads(row["values_json"])
                if not isinstance(values, list) or len(values) != 1:
                    raise ValueError("expected one value in values_json")
                value = values[0]
                if value is not None:
                    if isinstance(value, bool) or not isinstance(value, (int, float)):
                        raise ValueError("expected a numeric value or null")
                    value = float(value)
                    if not math.isfinite(value):
                        raise ValueError("non-finite value")
                sample = Sample(int(row["received_timestamp_ms"]),
                                int(row["source_timestamp_ms"]) if row["source_timestamp_ms"] else None,
                                row["sender_id"], row["id"], value)
                _ = sample.time  # Validate the timestamp before displaying/exporting.
                streams[kind].append(sample)
            except (ValueError, TypeError, OverflowError, OSError) as exc:
                raise ValueError(f"CSV line {line} ({kind}): {exc}") from exc
    streams = {key: sorted(rows, key=lambda row: row.received_ms)
               for key, rows in streams.items() if rows}
    if not streams:
        raise ValueError("No 1000kg load-cell or pressure-transducer data found.")
    return streams


def select_window(streams, start_ms, end_ms):
    if start_ms > end_ms:
        raise ValueError("Start must be before end.")
    selected = {key: [row for row in rows if start_ms <= row.received_ms <= end_ms]
                for key, rows in streams.items()}
    if not any(selected.values()):
        raise ValueError("No samples in this time window.")
    return selected


def extrema(rows):
    valid = [row for row in rows if row.value is not None]
    return (min(valid, key=lambda row: row.value), max(valid, key=lambda row: row.value)) if valid else None


def recalculate_kg(streams, slope=None, intercept=None):
    """Recover a consistent linear mapping per sender or apply supplied coefficients.

    Inference verifies all matched data, so clipped/nonlinear/calibration-changing
    recordings cannot silently produce a misleading least-squares mapping.
    """
    raw = streams.get("KG1000", [])
    if not raw:
        raise ValueError("Raw KG1000 readings are required for recalculation.")
    if (slope is None) != (intercept is None):
        raise ValueError("Supply both kg slope and intercept.")
    if slope is not None and not all(math.isfinite(v) for v in (slope, intercept)):
        raise ValueError("Mapping coefficients must be finite.")
    mappings = {}
    notes = []
    fits = []
    recorded = defaultdict(deque)
    for row in streams.get("LOADCELL_WEIGHT_KG", []):
        recorded[(row.received_ms, row.source_ms, row.sender)].append(row.value)
    pairs = defaultdict(list)
    for row in raw:
        matches = recorded[(row.received_ms, row.source_ms, row.sender)]
        if matches:
            value = matches.popleft()
            if row.value is not None and value is not None:
                pairs[row.sender].append((row.value, value))
    if slope is None:
        for sender in sorted({row.sender for row in raw}):
            points = pairs[sender]
            if len(points) < 3:
                raise ValueError(f"Insufficient matched raw/kg readings for {sender}; supply --kg-slope and --kg-intercept, or use --kg-mode recorded.")
            mean_x = math.fsum(x for x, _ in points) / len(points)
            mean_y = math.fsum(y for _, y in points) / len(points)
            variance = math.fsum((x - mean_x) ** 2 for x, _ in points)
            if variance == 0:
                raise ValueError(f"Raw values are constant for {sender}; supply an explicit mapping.")
            m = math.fsum((x - mean_x) * (y - mean_y) for x, y in points) / variance
            b = mean_y - m * mean_x
            error = max(abs(y - (m * x + b)) for x, y in points)
            tolerance = max(0.001, (max(y for _, y in points) - min(y for _, y in points)) * 1e-5)
            if error > tolerance:
                raise ValueError(f"{sender}: recorded kg is not a consistent linear mapping (max error {error:.6g} kg). Supply --kg-slope/--kg-intercept from the recording's calibration, or use --kg-mode recorded.")
            mappings[sender] = (m, b)
            notes.append((f"Kg mapping ({sender})", f"Inferred from {len(points)} timestamp-matched pairs: kg = {m:.15g} * raw + ({b:.15g}); max residual {error:.6g} kg"))
    else:
        for sender in {row.sender for row in raw}:
            mappings[sender] = (slope, intercept)
        notes.append(("Kg mapping", f"User supplied: kg = {slope:.15g} * raw + ({intercept:.15g})"))
    for sender, (m, b) in sorted(mappings.items()):
        points = pairs[sender]
        residuals = [y - (m * x + b) for x, y in points]
        sse = math.fsum(error ** 2 for error in residuals)
        mean_y = math.fsum(y for _, y in points) / len(points) if points else 0
        sst = math.fsum((y - mean_y) ** 2 for _, y in points)
        fits.append({"sender": sender, "method": "Inferred from recording" if slope is None else "User supplied",
                     "slope": m, "intercept": b, "points": sorted(points),
                     "r_squared": 1 - sse / sst if sst else None,
                     "rmse": math.sqrt(sse / len(points)) if points else None,
                     "max_error": max(map(abs, residuals)) if residuals else None})
    signed = []
    for row in raw:
        m, b = mappings[row.sender]
        value = None if row.value is None else m * row.value + b
        if value is not None and not math.isfinite(value):
            raise ValueError("Mapping produced a non-finite kg value.")
        signed.append(replace(row, value=value))
    result = dict(streams)
    result["RECALCULATED_KG"] = signed
    result["ABSOLUTE_KG"] = [replace(row, value=None if row.value is None else abs(row.value)) for row in signed]
    notes.append(("Absolute kg", "abs(recalculated signed kg); original signed readings retained. No clipping or baseline subtraction."))
    return result, notes, fits


def write_calibration(workbook, fits):
    """Save the exact mapping, diagnostics, all supporting pairs, and fit charts."""
    from openpyxl.chart import Reference, ScatterChart, Series

    summary = workbook.create_sheet("Calibration")
    summary.append(("Mapping", "signed kg = slope * raw + intercept; absolute kg = abs(signed kg)"))
    summary.append(("Scope", "Fit uses full recording, independent of the selected export window"))
    summary.append(("Interpretation", "Reconstructs the conversion applied during recording; not independent physical calibration"))
    summary.append(("Input units", "Recorded KG1000 raw units; do not assume these are actual ADC-input volts"))
    summary.append(("Diagnostics", "Residual = recorded kg minus mapped kg; R²/RMSE compare against recorded kg"))
    summary.append(("Sender", "Method", "Slope (kg/raw)", "Intercept (kg)", "Matched pairs",
                    "R squared", "RMSE (kg)", "Max abs residual (kg)", "Raw minimum", "Raw maximum", "Zero raw"))
    for index, fit in enumerate(fits, 1):
        points = fit["points"]
        m, b = fit["slope"], fit["intercept"]
        summary.append((fit["sender"], fit["method"], m, b, len(points), fit["r_squared"],
                        fit["rmse"], fit["max_error"], points[0][0] if points else None,
                        points[-1][0] if points else None, -b / m if m else None))
        summary.cell(summary.max_row, 1).data_type = "s"
        sheet = workbook.create_sheet(f"Calibration pairs {index}")
        sheet.append(("Raw", "Recorded kg", "Mapped signed kg", "Residual kg", None, "Fit raw", "Fit kg"))
        for x, y in points:
            mapped = m * x + b
            sheet.append((x, y, mapped, y - mapped))
        sheet.freeze_panes = "A2"
        for column in "ABCDEFG":
            sheet.column_dimensions[column].width = 24
        if not points:
            continue
        for number, x in enumerate((points[0][0], points[-1][0]), 2):
            sheet.cell(number, 6, x)
            sheet.cell(number, 7, m * x + b)
        chart = ScatterChart()
        chart.title = f"{fit['sender']}: kg = {m:.9g} × raw + ({b:.9g})"
        chart.x_axis.title = "Recorded raw units"
        chart.y_axis.title = "Signed kg"
        chart.width, chart.height = 28, 13
        observed = Series(Reference(sheet, min_col=2, min_row=2, max_row=len(points) + 1),
                          Reference(sheet, min_col=1, min_row=2, max_row=len(points) + 1), title="Recorded pairs")
        observed.marker.symbol = "circle"
        observed.marker.size = 3
        observed.graphicalProperties.line.noFill = True
        chart.series.append(observed)
        fitted = Series(Reference(sheet, min_col=7, min_row=2, max_row=3),
                        Reference(sheet, min_col=6, min_row=2, max_row=3), title="Applied linear mapping")
        fitted.marker.symbol = "none"
        fitted.graphicalProperties.line.solidFill = "C62828"
        chart.series.append(fitted)
        sheet.add_chart(chart, "I2")
    for column in "ABCDEFGHIJK":
        summary.column_dimensions[column].width = 25
    summary.column_dimensions["B"].width = 90
    summary.freeze_panes = "C7"


def export_excel(streams, path, source, start_ms, end_ms, notes=(), fits=()):
    from openpyxl import Workbook
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Font
    from openpyxl.utils.datetime import to_excel

    selected = select_window(streams, start_ms, end_ms)
    if any(len(rows) > 1048575 for rows in selected.values()):
        raise ValueError("Selection exceeds Excel's row limit; select a shorter window.")
    workbook = Workbook()
    info = workbook.active
    info.title = "Recording"
    for row in [
        ("Input CSV", str(source.resolve())),
        ("Chart clock", "Received timestamps, UTC; millisecond precision"),
        ("Selected start UTC", datetime.fromtimestamp(start_ms / 1000, timezone.utc).isoformat()),
        ("Selected end UTC", datetime.fromtimestamp(end_ms / 1000, timezone.utc).isoformat()),
        ("Values", "Recorded channels preserved; separately labelled recalculated/absolute kg when enabled. No interpolation."),
        ("Source clock", "Original device milliseconds; may differ between senders"),
        ("Missing values", "Recorded nulls remain blank; missing channels are not synthesized"),
    ]:
        info.append(row)
    for row in notes:
        info.append(row)
    info.append(("Data type", "Samples in selection"))
    charts = workbook.create_sheet("Charts")
    summary = workbook.create_sheet("Min and max")
    summary.append(("Channel", "Unit", "Minimum", "Minimum UTC", "Maximum", "Maximum UTC"))
    chart_index = 0
    for kind, rows in selected.items():
        info.append((kind, len(rows)))
        if not rows:
            continue
        title, unit = CHANNELS[kind]
        sheet = workbook.create_sheet(title)
        sheet.append(("Received time (UTC)", "Received timestamp (ms)",
                      "Source timestamp (ms)", "Sender", "CSV row ID", f"Value ({unit})"))
        for number, row in enumerate(rows, 2):
            sheet.append((row.time.replace(tzinfo=None), row.received_ms,
                          row.source_ms, row.sender, row.row_id, row.value))
            # Treat CSV strings as literal text, even if they start with '='.
            for column in (4, 5):
                sheet.cell(number, column).data_type = "s"
        for cell in sheet["A"][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss.000"
        for column in ("B", "C"):
            for cell in sheet[column][1:]:
                cell.number_format = "0"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column, width in zip("ABCDEF", (28, 25, 25, 15, 15, 24)):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        chart = ScatterChart()
        chart.title = title
        limits = extrema(rows)
        if limits:
            low, high = limits
            chart.title = f"{title} | Min {low.value:.6g} / Max {high.value:.6g} {unit}"
            summary.append((title, unit, low.value, low.time.replace(tzinfo=None),
                            high.value, high.time.replace(tzinfo=None)))
            for column in (4, 6):
                summary.cell(summary.max_row, column).number_format = "yyyy-mm-dd hh:mm:ss.000"
        chart.x_axis.title = "Received time (UTC)"
        chart.x_axis.numFmt = "hh:mm:ss.000"
        chart.y_axis.title = unit
        chart.scatterStyle = "line"
        chart.display_blanks = "gap"
        chart.width, chart.height = 28, 12
        chart.legend = None
        chart.x_axis.scaling.min = to_excel(datetime.fromtimestamp(start_ms / 1000, timezone.utc).replace(tzinfo=None))
        chart.x_axis.scaling.max = to_excel(datetime.fromtimestamp(end_ms / 1000, timezone.utc).replace(tzinfo=None)) if end_ms > start_ms else None
        # Separate series per sender so unrelated devices are never joined.
        senders = sorted({row.sender for row in rows})
        for index, sender in enumerate(senders):
            column = 7 + index
            sheet.cell(1, column, f"Chart values: {sender}").data_type = "s"
            for number, row in enumerate(rows, 2):
                if row.sender == sender:
                    sheet.cell(number, column, row.value)
            series = Series(Reference(sheet, min_col=column, min_row=2, max_row=len(rows) + 1),
                            Reference(sheet, min_col=1, min_row=2, max_row=len(rows) + 1), title=sender)
            series.marker.symbol = "none"
            chart.series.append(series)
        if len(senders) > 1:
            from openpyxl.chart.legend import Legend
            chart.legend = Legend()
        if limits:
            for index, (name, row, color) in enumerate(zip(("Min", "Max"), limits, ("1565C0", "C62828"))):
                column = 7 + len(senders) + index * 2
                sheet.cell(1, column, f"{name} UTC")
                sheet.cell(1, column + 1, f"{name} ({unit})")
                sheet.cell(2, column, row.time.replace(tzinfo=None)).number_format = "yyyy-mm-dd hh:mm:ss.000"
                sheet.cell(2, column + 1, row.value)
                series = Series(Reference(sheet, min_col=column + 1, min_row=2, max_row=2),
                                Reference(sheet, min_col=column, min_row=2, max_row=2), title=name)
                series.marker.symbol = "circle"
                series.marker.size = 8
                series.marker.graphicalProperties.solidFill = color
                series.graphicalProperties.line.noFill = True
                series.dLbls = DataLabelList()
                series.dLbls.showVal = True
                series.dLbls.dLblPos = "b" if name == "Min" else "t"
                chart.series.append(series)
        charts.add_chart(chart, f"A{1 + chart_index * 25}")
        chart_index += 1
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 110
    for column in "ABCDEF":
        summary.column_dimensions[column].width = 28
    summary.freeze_panes = "A2"
    if fits:
        if any(len(fit["points"]) > 1048575 for fit in fits):
            raise ValueError("Calibration pairs exceed Excel's row limit.")
        write_calibration(workbook, fits)
    path.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation prevents accidental replacement of an earlier export.
    with path.open("xb") as output:
        workbook.save(output)
    return sum(len(rows) for rows in selected.values())


def show_selector(streams, path, source, notes=(), kg_mode="recorded", fits=()):
    import matplotlib.dates as dates
    import matplotlib.pyplot as plt
    from matplotlib.widgets import Button, SpanSelector

    kg_channel = {"recorded": "LOADCELL_WEIGHT_KG", "signed": "RECALCULATED_KG", "absolute": "ABSOLUTE_KG"}[kg_mode]
    displayed = {key: rows for key, rows in streams.items()
                 if key not in ("LOADCELL_WEIGHT_KG", "RECALCULATED_KG", "ABSOLUTE_KG") or key == kg_channel}
    # Place the selected kg representation immediately after raw loadcell values.
    displayed = dict(sorted(displayed.items(), key=lambda item: 0 if item[0] == "KG1000" else 1 if item[0] == kg_channel else 2))
    fig, axes = plt.subplots(len(displayed), 1, sharex=True, figsize=(14, 9), squeeze=False)
    fig.subplots_adjust(left=0.10, right=0.97, top=0.91, bottom=0.18, hspace=0.15)
    fig.suptitle("1000kg load cell and pressure transducer — drag across a plot to select a window")
    axes = axes[:, 0]
    annotations = []
    for axis, (kind, rows) in zip(axes, displayed.items()):
        for sender in sorted({row.sender for row in rows}):
            samples = [row for row in rows if row.sender == sender]
            axis.plot([row.time for row in samples],
                      [row.value if row.value is not None else math.nan for row in samples],
                      linewidth=0.8, label=sender)
        axis.set_ylabel(f"{CHANNELS[kind][0]}\n({CHANNELS[kind][1]})")
        axis.grid(alpha=0.25)
        axis.legend(loc="upper right")
        label = axis.text(0.01, 0.94, "", transform=axis.transAxes, va="top", fontsize=9,
                          bbox={"facecolor": "white", "alpha": 0.85, "edgecolor": "none"})
        markers = axis.scatter([math.nan, math.nan], [math.nan, math.nan],
                               c=["tab:blue", "tab:red"], s=35, zorder=5)
        annotations.append((label, markers))

    def update_extrema(start=None, end=None):
        for (kind, rows), (label, markers) in zip(displayed.items(), annotations):
            selected = rows if start is None else [row for row in rows if start <= row.received_ms <= end]
            limits = extrema(selected)
            if limits:
                low, high = limits
                unit = CHANNELS[kind][1]
                scope = "Selection" if start is not None else "Full recording"
                label.set_text(f"{scope}: Min {low.value:.6g} / Max {high.value:.6g} {unit}")
                markers.set_offsets([(dates.date2num(row.time), row.value) for row in limits])
                markers.set_visible(True)
            else:
                label.set_text("No valid samples in selection")
                markers.set_visible(False)

    update_extrema()
    axes[-1].xaxis.set_major_formatter(dates.DateFormatter("%Y-%m-%d\n%H:%M:%S", tz=timezone.utc))
    axes[-1].set_xlabel("Received time (UTC)")
    # Lock the data limits before selectors add their initially zero-width artists.
    # Otherwise those artists can autoscale a datetime axis back to the Unix epoch.
    left = min(dates.date2num(rows[0].time) for rows in streams.values())
    right = max(dates.date2num(rows[-1].time) for rows in streams.values())
    padding = max((right - left) * 0.01, 1 / 86400000)
    axes[-1].set_xlim(left - padding, right + padding)
    status = fig.text(0.10, 0.075, "Drag to choose a window. Close the window to cancel.", fontsize=10)
    selection = []
    selectors = []

    def on_select(low, high):
        selection[:] = [round(dates.num2date(value).timestamp() * 1000) for value in (low, high)]
        for selector in selectors:
            selector.extents = (low, high)
            selector.set_visible(True)
        count = sum(sum(selection[0] <= row.received_ms <= selection[1] for row in rows)
                    for rows in streams.values())
        bounds = [datetime.fromtimestamp(value / 1000, timezone.utc).isoformat(timespec="milliseconds")
                  for value in selection]
        status.set_text(f"{bounds[0]} to {bounds[1]} | {count:,} readings")
        update_extrema(*selection)
        fig.canvas.draw_idle()

    for axis in axes:
        selectors.append(SpanSelector(axis, on_select, "horizontal", useblit=True,
                                      props={"facecolor": "tab:orange", "alpha": 0.25},
                                      interactive=True, drag_from_anywhere=True))
    button = Button(fig.add_axes([0.73, 0.015, 0.23, 0.045]), "Export selection")

    def on_export(_event):
        if not selection:
            status.set_text("Select a time window by dragging across a plot first.")
        else:
            try:
                count = export_excel(streams, path, source, *selection, notes=notes, fits=fits)
            except (ValueError, OSError) as exc:
                status.set_text(f"Export failed: {exc}")
            else:
                print(f"Exported {count:,} readings to {path.resolve()}", flush=True)
                plt.close(fig)
        fig.canvas.draw_idle()

    button.on_clicked(on_export)
    plt.show()


def parse_time(value):
    try:
        date = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if date.tzinfo is None:
            date = date.replace(tzinfo=timezone.utc)
        return round(date.timestamp() * 1000)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("Use an ISO timestamp such as 2026-09-19T09:05:00Z") from exc


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("csv", type=Path)
    parser.add_argument("-o", "--output", type=Path, help="Default: <recording>_loadcell_pt_<kg-mode>.xlsx in the current directory")
    parser.add_argument("--start", type=parse_time, help="Inclusive start timestamp (skips the selection GUI)")
    parser.add_argument("--end", type=parse_time, help="Inclusive end timestamp (skips the selection GUI)")
    parser.add_argument("--all", action="store_true", help="Export the entire recording without a GUI")
    parser.add_argument("--kg-mode", choices=("absolute", "signed", "recorded"), default="absolute",
                        help="Kg plot mode (default: absolute recalculated kg); Excel retains signed values too")
    parser.add_argument("--kg-slope", type=float, help="Known raw-to-kg slope; otherwise infer from this recording")
    parser.add_argument("--kg-intercept", type=float, help="Known raw-to-kg intercept including the recording's tare")
    args = parser.parse_args()
    if (args.start is None) != (args.end is None):
        parser.error("--start and --end must be supplied together")
    if args.all and args.start is not None:
        parser.error("Use --all or --start/--end")
    if (args.kg_slope is None) != (args.kg_intercept is None):
        parser.error("--kg-slope and --kg-intercept must be supplied together")
    if args.kg_mode == "recorded" and args.kg_slope is not None:
        parser.error("Explicit kg mapping requires --kg-mode signed or absolute")
    output = args.output or Path(f"{args.csv.stem}_loadcell_pt_{args.kg_mode}.xlsx")
    if output.suffix.lower() != ".xlsx":
        parser.error("Output must have an .xlsx extension")
    if output.exists():
        parser.error(f"Output already exists: {output}; choose another --output path")
    streams = read_recording(args.csv)
    notes = [("Kg plot mode", args.kg_mode)]
    fits = []
    if args.kg_mode != "recorded":
        streams, calibration_notes, fits = recalculate_kg(streams, args.kg_slope, args.kg_intercept)
        notes.extend(calibration_notes)
        for name, note in calibration_notes:
            print(f"{name}: {note}")
    for kind, rows in streams.items():
        print(f"{kind}: {len(rows):,} readings ({rows[0].time.isoformat()} to {rows[-1].time.isoformat()})")
    for kind in ("KG1000", "LOADCELL_WEIGHT_KG", "PRESSURE_TRANSDUCER_CALIBRATED"):
        if kind not in streams:
            print(f"Warning: {kind} is absent from the recording.", file=sys.stderr)
    if not any(kind in streams for kind in ("FUEL_TANK_PRESSURE", "IADC")):
        print("Warning: no raw PT channel in the recording.", file=sys.stderr)
    if args.all or args.start is not None:
        start = min(rows[0].received_ms for rows in streams.values()) if args.all else args.start
        end = max(rows[-1].received_ms for rows in streams.values()) if args.all else args.end
        count = export_excel(streams, output, args.csv, start, end, notes=notes, fits=fits)
        print(f"Exported {count:,} readings to {output.resolve()}")
    else:
        import openpyxl  # Check the export dependency before opening the GUI.
        show_selector(streams, output, args.csv, notes=notes, kg_mode=args.kg_mode, fits=fits)


if __name__ == "__main__":
    try:
        main()
    except ImportError as exc:
        sys.exit(f"Missing dependency: {exc}. Install with: python3 -m pip install matplotlib openpyxl")
    except (ValueError, OSError) as exc:
        sys.exit(f"Error: {exc}")
    except KeyboardInterrupt:
        sys.exit("Cancelled.")
